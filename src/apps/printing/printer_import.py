"""Excel import planning for the existing Printing models; no schema changes."""
from collections import Counter, defaultdict
from copy import copy
from dataclasses import dataclass, field
from ipaddress import ip_address
from pathlib import Path
import re
from zipfile import BadZipFile

from django.core.exceptions import ValidationError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from apps.accounts.models import Branch, User
from apps.inventory.models import OrganizationalLocation
from .models import PrintingDevice


SHEET_NAME = "IMPORTACION_PRINTING"
COLUMNS = (
    "N° origen", "ID origen", "Serie", "Modelo", "Sede", "Dependencia",
    "Responsable", "IP anterior", "IP actual", "Conexión", "Calcomanía",
    "Nombre impresora", "Contrato", "Tipo servicio", "Estado validación",
    "Observaciones",
)
START_NOTES = "[[NOVADESK_IMPORT_PRINTERS_V1]]"
END_NOTES = "[[/NOVADESK_IMPORT_PRINTERS_V1]]"
MAX_ROWS = 50000
BRANCH_ALIASES = {
    "villa elisa": "PLANTA-VILLA-ELISA",
    "viila elisa": "PLANTA-VILLA-ELISA",
    "centro": "OFICINA-CENTRAL",
    "troche": "PLANTA-MAURICIO-JOSE-TROCHE",
}
PROVISIONAL_LOCATION_CODE = "UBICACION-PENDIENTE"
LEXMARK_MODEL = re.compile(r"^(?:LEXMARK\s+)?(?:MX622|MX611|MX521|MX632|MX711|MX722|CX622|CX725)(?:[A-Z]+)?$", re.IGNORECASE)



class PrinterImportError(Exception):
    """A critical file/import error, requiring no commit or full rollback."""


def normalize_text(value):
    if value is None:
        return ""
    return " ".join(str(value).split())


def serial_key(value):
    return normalize_text(value).casefold()


def validate_serial(value):
    if not isinstance(value, str):
        raise ValidationError("Serie debe ser texto; no se pueden recuperar ceros iniciales de celdas numéricas.")
    value = normalize_text(value)
    if not value or len(value) > 150 or not any(char.isalnum() for char in value):
        raise ValidationError("Serie ausente, inválida o superior a 150 caracteres.")
    if value.casefold() in {"n/a", "na", "s/n", "sn", "sin serie", "sin dato", "sin datos", "pendiente", "null", "none", "0"}:
        raise ValidationError("Serie contiene un marcador sin identificación válida.")
    return value


def read_import_rows(filename):
    path = Path(filename)
    if path.suffix.lower() != ".xlsx":
        raise PrinterImportError("El archivo debe tener extensión .xlsx.")
    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        if SHEET_NAME not in workbook.sheetnames:
            raise PrinterImportError(f"Falta la hoja {SHEET_NAME}.")
        sheet = workbook[SHEET_NAME]
        # Use actual row data even when a producer wrote incorrect dimensions.
        sheet.reset_dimensions()
        iterator = sheet.iter_rows()
        header = next(iterator, ())
        headings = [normalize_text(cell.value) for cell in header]
        normalized = [name.casefold() for name in headings]
        if len([name for name in normalized if name]) != len(set(name for name in normalized if name)):
            raise PrinterImportError("Hay encabezados duplicados.")
        missing = [name for name in COLUMNS if name.casefold() not in normalized]
        if missing:
            raise PrinterImportError("Faltan columnas: " + ", ".join(missing))
        indices = {name: normalized.index(name.casefold()) for name in COLUMNS}
        rows = []
        for number, cells in enumerate(iterator, 2):
            if number > MAX_ROWS + 1:
                raise PrinterImportError(f"El archivo supera el límite de {MAX_ROWS} filas.")
            if not any(normalize_text(cell.value) for cell in cells):
                continue
            row = {name: cells[index].value if index < len(cells) else None for name, index in indices.items()}
            formulas = [name for name, index in indices.items() if name != "Contrato" and index < len(cells) and cells[index].data_type == "f"]
            rows.append((number, row, formulas))
        return rows
    except (OSError, BadZipFile, InvalidFileException, ValueError, KeyError) as exc:
        raise PrinterImportError("No se pudo leer el Excel. Verifique ruta, formato y archivo.") from exc
    finally:
        if workbook is not None:
            workbook.close()


def _unique_match(objects, value, attributes):
    key = serial_key(value)
    matches = [obj for obj in objects if any(serial_key(getattr(obj, attribute)) == key for attribute in attributes)]
    return matches[0] if len(matches) == 1 else None


def _import_notes(existing, row, pending):
    existing = existing or ""
    if existing.count(START_NOTES) != existing.count(END_NOTES) or existing.count(START_NOTES) > 1:
        raise ValidationError("El bloque de notas de importación existente está dañado; revisar manualmente.")
    labels = ("N° origen", "ID origen", "Sede", "Dependencia", "Responsable", "IP anterior", "IP actual", "Conexión", "Calcomanía", "Nombre impresora", "Tipo servicio", "Estado validación", "Observaciones")
    lines = [f"{name}: {row[name]}" for name in labels if row[name]]
    if pending:
        lines.append("Datos pendientes: " + "; ".join(pending))
    block = "\n".join([START_NOTES, *lines, END_NOTES])
    if START_NOTES in existing:
        start = existing.index(START_NOTES)
        end = existing.index(END_NOTES) + len(END_NOTES)
        if end <= start:
            raise ValidationError("Orden inválido de marcadores de notas.")
        return existing[:start] + block + existing[end:]
    return existing + ("\n\n" if existing else "") + block


@dataclass
class ImportRow:
    number: int
    status: str
    device: PrintingDevice | None = None
    fields: list = field(default_factory=list)
    pending: list = field(default_factory=list)
    errors: list = field(default_factory=list)


class PrinterImportPlanner:
    def __init__(self, *, brand="", lock=False):
        self.brand = normalize_text(brand)
        devices = PrintingDevice.objects.select_related("asset", "branch", "organizational_location")
        if lock:
            # Lock existing devices, not nullable joins (important on PostgreSQL).
            list(PrintingDevice.objects.select_for_update().values_list("pk", flat=True))
        self.devices = defaultdict(list)
        for device in devices:
            key = serial_key(device.effective_serial_number)
            if key:
                self.devices[key].append(device)
        self.branches = list(Branch.objects.filter(is_active=True))
        self.locations = list(OrganizationalLocation.objects.filter(is_active=True, branch__is_active=True).select_related("branch", "parent"))
        self.users = list(User.objects.filter(is_active=True, approval_status=User.ApprovalStatus.APPROVED))

    def plan(self, rows):
        serial_counts = Counter(serial_key(row["Serie"]) for _, row, _ in rows if serial_key(row["Estado validación"]) != "revisar" and normalize_text(row["Serie"]))
        return [self.plan_row(number, row, formulas, serial_counts) for number, row, formulas in rows]

    def plan_row(self, number, raw, formulas, serial_counts):
        if serial_key(raw["Estado validación"]) == "revisar":
            return ImportRow(number, "IGNORADO")
        try:
            if formulas:
                raise ValidationError("La fila contiene fórmulas; se requieren valores literales.")
            serial = validate_serial(raw["Serie"])
            key = serial_key(serial)
            if serial_counts[key] > 1:
                raise ValidationError("Serie repetida en el Excel: se rechazan todas sus filas, sin elegir una.")
            existing = self.devices.get(key, [])
            if len(existing) > 1:
                raise ValidationError("La serie coincide con varios equipos existentes; resolver duplicados antes de importar.")
            row = {name: normalize_text(raw[name]) for name in COLUMNS}
            for name in ("IP anterior", "IP actual"):
                if row[name]:
                    try:
                        row[name] = str(ip_address(row[name]))
                    except ValueError as exc:
                        raise ValidationError(f"{name}: dirección IP inválida.") from exc
            new = not existing
            original = existing[0] if existing else None
            device = copy(original) if original else PrintingDevice()
            if original:
                device._state = copy(original._state)
                device._state.fields_cache = original._state.fields_cache.copy()
            changes = {}
            pending = []
            if new or not normalize_text(device.serial_number):
                changes["serial_number"] = serial
            elif device.serial_number != normalize_text(device.serial_number):
                changes["serial_number"] = normalize_text(device.serial_number)
            if row["Modelo"]:
                changes["model"] = row["Modelo"]
            if not LEXMARK_MODEL.fullmatch(row["Modelo"]):
                raise ValidationError("REVISAR: modelo/marca no reconocido; no se infiere marca.")
            if self.brand and self.brand.casefold() != "lexmark":
                raise ValidationError("REVISAR: marca indicada incompatible con la familia Lexmark confirmada.")
            if device.brand and device.brand.casefold() != "lexmark":
                raise ValidationError("REVISAR: marca existente incompatible con el modelo Lexmark.")
            changes["brand"] = "Lexmark"
            branch = device.branch
            if row["Sede"]:
                branch_value = BRANCH_ALIASES.get(serial_key(row["Sede"]), row["Sede"])
                branch = _unique_match(self.branches, branch_value, ("code", "name"))
                if branch:
                    changes["branch"] = branch
                else:
                    pending.append("Sede sin equivalencia inequívoca")
            if row["Dependencia"] and branch:
                location = _unique_match([item for item in self.locations if item.branch_id == branch.pk], row["Dependencia"], ("code", "name", "full_path"))
                if location:
                    changes["organizational_location"] = location
                else:
                    pending.append("Dependencia sin ubicación equivalente dentro de la sede")
            elif row["Dependencia"]:
                pending.append("Dependencia sin sede identificada")
            if branch and "organizational_location" not in changes:
                current_location = device.organizational_location
                if not current_location or current_location.branch_id != branch.pk or not current_location.is_active:
                    provisional = _unique_match(
                        [item for item in self.locations if item.branch_id == branch.pk],
                        PROVISIONAL_LOCATION_CODE, ("code",),
                    )
                    if provisional:
                        changes["organizational_location"] = provisional
                        pending.append("Ubicaci\u00f3n provisional pendiente de verificar")
            if row["Responsable"]:
                responsible = _unique_match(self.users, row["Responsable"], ("email", "username"))
                if responsible:
                    changes["responsible_user"] = responsible
                else:
                    pending.append("Responsable sin correo/username local inequívoco")
            if row["IP actual"] or row["IP anterior"]:
                pending.append("IP conservada como texto; no constituye detección de red")
            for name, value in changes.items():
                setattr(device, name, value)
            notes_row = row.copy()
            # Keep the complete original dependence text, including whitespace.
            notes_row["Dependencia"] = "" if raw["Dependencia"] is None else str(raw["Dependencia"])
            device.notes = _import_notes(device.notes, notes_row, pending)
            changes["notes"] = device.notes
            device.full_clean()
            fields = list(changes) if new else [name for name in changes if getattr(original, name) != getattr(device, name)]
            status = "NUEVO" if new else "ACTUALIZADO" if fields else "SIN CAMBIOS"
            return ImportRow(number, status, device, fields, pending=pending)
        except ValidationError as exc:
            return ImportRow(number, "RECHAZADO", errors=exc.messages)


def save_import_row(row):
    """Called only inside the command's single transaction.atomic block."""
    if row.status == "NUEVO":
        row.device.save()
    elif row.status == "ACTUALIZADO" and row.fields:
        row.device.save(update_fields=[*row.fields, "updated_at"])
